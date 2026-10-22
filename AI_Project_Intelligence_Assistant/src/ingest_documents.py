import io
import hashlib

from pypdf import PdfReader
from pptx import Presentation
from docx import Document
from openpyxl import load_workbook

from src.graph_client import (
    list_onedrive_files,
    list_folder_items,
    download_file_content,
)


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".docx",
    ".xlsx",
    ".pdf",
    ".pptx",
    ".csv",
}


def extract_text(file_name, content):
    name = file_name.lower()

    # TXT and Markdown files
    if name.endswith((".txt", ".md")):
        return content.decode("utf-8-sig", errors="replace")

    # Word documents
    if name.endswith(".docx"):
        document = Document(io.BytesIO(content))
        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts)

    # PDF documents
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        parts = []

        for page in reader.pages:
            text = page.extract_text()
            if text:
                parts.append(text)

        return "\n".join(parts)

    # PowerPoint presentations
    if name.endswith(".pptx"):
        presentation = Presentation(io.BytesIO(content))
        parts = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text)

                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(
                                cell.text for cell in row.cells
                            )
                        )

        return "\n".join(parts)

    # CSV files
    if name.endswith(".csv"):
        return content.decode("utf-8-sig", errors="replace")

    # Excel workbooks
    if name.endswith(".xlsx"):
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
        parts = []

        try:
            for sheet in workbook.worksheets:
                parts.append(f"Sheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value) if value is not None else ""
                        for value in row
                    ]

                    if any(value.strip() for value in values):
                        parts.append(" | ".join(values))
        finally:
            workbook.close()

        return "\n".join(parts)

    return ""


def discover_files(folder_id, folder_path="Dataset"):
    files = []

    for item in list_folder_items(folder_id):
        name = item.get("name", "")
        item_path = f"{folder_path}/{name}"

        if "folder" in item:
            files.extend(
                discover_files(item["id"], item_path)
            )

        elif any(
            name.lower().endswith(ext)
            for ext in SUPPORTED_EXTENSIONS
        ):
            files.append({
                "id": item["id"],
                "name": name,
                "path": item_path,
            })

    return files


def load_dataset():
    root_items = list_onedrive_files()

    dataset_folder = next(
        (
            item
            for item in root_items
            if item.get("name", "").lower() == "dataset"
            and "folder" in item
        ),
        None,
    )

    if dataset_folder is None:
        raise FileNotFoundError(
            "Could not find a Dataset folder in the OneDrive root."
        )

    files = discover_files(dataset_folder["id"])
    documents = []

    for file in files:
        print(f"Reading: {file['path']}")

        content = download_file_content(file["id"])
        text = extract_text(file["name"], content).strip()

        if not text:
            print("  Skipped: no extractable text.")
            continue

        documents.append({
            "id": hashlib.sha256(file["id"].encode("utf-8")).hexdigest(),
            "title": file["name"],
            "source": file["path"],
            "content": text,
        })

    return documents


if __name__ == "__main__":
    documents = load_dataset()

    print(
        f"\nSuccessfully extracted {len(documents)} documents."
    )

    for document in documents:
        print(
            f"- {document['source']}: "
            f"{len(document['content'])} characters"
        )
